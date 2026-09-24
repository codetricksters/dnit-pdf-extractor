"""Gera as fixtures do teste ponta a ponta (uso único).

    uv run python scripts/gerar_fixtures_e2e.py

É o único código do projeto que lê a planilha de referência
``tmp/Reequilíbrio - 26 - Contrato 716-22.xlsx``, e só para produzir as fixtures
versionadas em ``tests/fixtures/``. O seed e a aplicação nunca a leem.

* ``igp_di.xlsx``: o IGP-DI de jan/2022 (a base, célula I6 do oráculo) e os meses
  da aba ``IGP - DI``, no formato do template de importação;
* ``delta_p_referencia.csv``: ``mes, delta_cap, delta_emul`` do oráculo,
  calculados pelo Excel, de forma independente da aplicação;
* ``contrato_ficticio.json``: itens reais de jan a dez/2023 extraídos de outro
  contrato, sob um cabeçalho fictício com a Data Base do oráculo.
"""

import csv
import json
import sys
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app import config  # noqa: E402
from app.services.importadores import gerar_template_igp_di  # noqa: E402

PLANILHA = config.PROJECT_ROOT / "tmp" / "Reequilíbrio - 26 - Contrato 716-22.xlsx"
RESULTADOS = config.PROJECT_ROOT / "data" / "jobs" / "37135d545e7a" / "results"
MEDICOES = [f"{n}ª MP_resultado.json" for n in range(10, 22)]  # jan–dez/2023
DESTINO = config.PROJECT_ROOT / "tests" / "fixtures"

ABA_ORACULO = "CÁLCULO DA VARIAÇÃO DE PREÇOS"
ABA_IGP = "IGP - DI"
LINHA_IGP = 26  # "IGP - DI", base ago/1994 = 100
BASE_IGP = (date(2022, 1, 1), "I6")

CABECALHO_FICTICIO = {
    "Contrato": "99 99999/2099 - CONSTRUTORA FICTÍCIA LTDA",
    "Data Base": "01/01/2022",
    "Número do Processo": "99999.999999/2099-99",
}


def _mes(valor) -> date:
    return (valor.date() if isinstance(valor, datetime) else valor).replace(day=1)


def _numero(valor) -> bool:
    return isinstance(valor, (int, float)) and not isinstance(valor, bool)


def igp_di(livro) -> dict[date, Decimal]:
    ws = livro[ABA_IGP]
    rotulo = ws.cell(LINHA_IGP, 1).value
    if not isinstance(rotulo, str) or "IGP" not in rotulo:
        raise SystemExit(f"A linha {LINHA_IGP} da aba '{ABA_IGP}' não é o IGP-DI: {rotulo!r}")
    valores = {}
    for celula in ws[3][2:]:
        if isinstance(celula.value, datetime):
            valor = ws.cell(LINHA_IGP, celula.column).value
            if _numero(valor):
                valores[_mes(celula.value)] = Decimal(str(valor))
    mes_base, endereco = BASE_IGP
    valores[mes_base] = Decimal(str(livro[ABA_ORACULO][endereco].value))
    return valores


def oraculo(livro, igp: dict[date, Decimal]) -> list[dict]:
    ws = livro[ABA_ORACULO]
    linhas = []
    for r in range(7, ws.max_row + 1):
        mes, cap, igp_mes, emul = (ws.cell(r, c).value for c in (8, 4, 9, 10))
        if not (_numero(cap) and _numero(emul)):
            break
        mes = _mes(mes)
        # A coluna I do oráculo e a aba IGP - DI têm de ser a mesma série.
        if Decimal(str(igp_mes)) != igp.get(mes):
            raise SystemExit(f"IGP-DI de {mes:%m/%Y} diverge entre as abas: {igp_mes} x {igp.get(mes)}")
        linhas.append({"mes": mes.isoformat(), "delta_cap": repr(cap), "delta_emul": repr(emul)})
    return linhas


def contrato() -> dict:
    arquivos = []
    for nome in MEDICOES:
        dados = json.loads((RESULTADOS / nome).read_text(encoding="utf-8"))
        arquivos.append({"arquivo": nome.replace("_resultado.json", ".pdf"), "rows": dados["rows"]})
    return {"header": CABECALHO_FICTICIO, "arquivos": arquivos}


def main() -> None:
    DESTINO.mkdir(parents=True, exist_ok=True)
    livro = openpyxl.load_workbook(PLANILHA, data_only=True)

    igp = igp_di(livro)
    (DESTINO / "igp_di.xlsx").write_bytes(
        gerar_template_igp_di([{"mes_ref": m, "valor": v} for m, v in igp.items()])
    )

    linhas = oraculo(livro, igp)
    with (DESTINO / "delta_p_referencia.csv").open("w", newline="", encoding="utf-8") as f:
        escritor = csv.DictWriter(f, fieldnames=["mes", "delta_cap", "delta_emul"])
        escritor.writeheader()
        escritor.writerows(linhas)

    (DESTINO / "contrato_ficticio.json").write_text(
        json.dumps(contrato(), ensure_ascii=False, indent=1), encoding="utf-8"
    )
    print(f"igp_di.xlsx: {len(igp)} meses; delta_p_referencia.csv: {len(linhas)} meses; "
          f"contrato_ficticio.json: {len(MEDICOES)} medições")


if __name__ == "__main__":
    main()

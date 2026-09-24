"""Leitura dos arquivos de índices: layout, validação por linha, tudo ou nada."""

import io
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from app.services import importadores
from app.services.delta_p import ANP_PRODUTO_CAP, BASE_IGP_DI, INDICE_IGP_DI
from app.services.importadores import ArquivoInvalido

from .indices_factory import linha_anp, linhas_anp, mes_igp, semana

OFICIAL = Path("tests/fixtures/anp_semanal.xls")
CAP = ANP_PRODUTO_CAP
PRECOS = [3.1, 4.02073, "***", 3.45, 3.3, 3.9]


# --- ANP ---------------------------------------------------------------------

def test_linha_vira_seis_registros_com_sem_cotacao_nulo():
    leitura = importadores.ler_linhas_anp(linhas_anp([linha_anp(CAP, date(2021, 12, 13), PRECOS)]))
    assert len(leitura.registros) == 6
    por_regiao = {r["regiao"]: r for r in leitura.registros}
    assert por_regiao["Nordeste"]["preco"] == Decimal("4.02073")
    assert por_regiao["Centro-Oeste"]["preco"] is None
    assert por_regiao["Brasil"]["preco"] == Decimal("3.90000")
    assert por_regiao["Nordeste"]["vigencia_inicio"] == date(2021, 12, 13)
    assert por_regiao["Nordeste"]["vigencia_fim"] == date(2021, 12, 19)
    assert leitura.avisos == []


def test_preco_e_quantizado_em_cinco_casas():
    precos = [1.2935999999999999] * 6
    leitura = importadores.ler_linhas_anp(linhas_anp([linha_anp(CAP, date(2013, 1, 1), precos)]))
    assert {r["preco"] for r in leitura.registros} == {Decimal("1.29360")}


def test_rodape_encerra_a_leitura():
    linhas = linhas_anp([linha_anp(CAP, date(2023, 1, 9), PRECOS)])
    linhas.append([CAP, "texto", "", 1, 1, 1, 1, 1, 1])  # depois do rodapé: ignorada
    assert len(importadores.ler_linhas_anp(linhas).registros) == 6


@pytest.mark.parametrize(
    "linha,trecho",
    [
        (linha_anp(CAP, date(2023, 1, 9), ["abc", 1, 1, 1, 1, 1]), "linha 10, Norte"),
        (linha_anp(CAP, date(2023, 1, 9), [1, -2, 1, 1, 1, 1]), "linha 10, Nordeste"),
        (linha_anp(CAP, date(2023, 1, 9), PRECOS, dias=-3), "linha 10: a semana termina"),
    ],
)
def test_linha_invalida_recusa_o_arquivo_inteiro(linha, trecho):
    boa = linha_anp(CAP, date(2023, 1, 16), PRECOS)
    with pytest.raises(ArquivoInvalido) as erro:
        importadores.ler_linhas_anp(linhas_anp([linha, boa]))
    assert "nada foi gravado" in erro.value.mensagem
    assert any(trecho in e for e in erro.value.erros)


def test_layout_errado_e_recusado():
    linhas = linhas_anp([linha_anp(CAP, date(2023, 1, 9), PRECOS)])
    linhas[8][4] = "Nordestee"
    with pytest.raises(ArquivoInvalido) as erro:
        importadores.ler_linhas_anp(linhas)
    assert "não é o arquivo de preços semanais da ANP" in erro.value.mensagem


def test_bytes_que_nao_sao_xls_sao_recusados():
    with pytest.raises(ArquivoInvalido) as erro:
        importadores.ler_anp(b"isto nao e uma planilha")
    assert "não é o arquivo de preços semanais da ANP" in erro.value.mensagem


def test_arquivo_sem_semanas_e_recusado():
    with pytest.raises(ArquivoInvalido):
        importadores.ler_linhas_anp(linhas_anp([]))


def test_produto_com_semana_repetida_e_pulado_com_aviso():
    glp = "Gás Liquefeito de Petróleo - GLP (R$/kg)"
    linhas = linhas_anp(
        [
            linha_anp(glp, date(2023, 1, 9), PRECOS),
            linha_anp(glp, date(2023, 1, 9), PRECOS),
            linha_anp(CAP, date(2023, 1, 9), PRECOS),
        ]
    )
    leitura = importadores.ler_linhas_anp(linhas)
    assert {r["produto"] for r in leitura.registros} == {CAP}
    (aviso,) = leitura.avisos
    assert glp in aviso and "não foi importado" in aviso


def test_limitar_resume_o_excesso():
    erros = [f"linha {i}" for i in range(80)]
    limitados = importadores.limitar(erros)
    assert len(limitados) == importadores.MAX_ERROS + 1
    assert limitados[-1] == "… e mais 30 erro(s)."


def test_arquivo_oficial():
    leitura = importadores.ler_anp(OFICIAL.read_bytes())
    assert len(leitura.registros) == 60114
    assert len({r["produto"] for r in leitura.registros}) == 20
    assert len(leitura.avisos) == 1 and "GLP" in leitura.avisos[0]
    chave = {(r["produto"], r["regiao"], r["vigencia_inicio"]): r["preco"] for r in leitura.registros}
    assert chave[(CAP, "Nordeste", date(2021, 12, 13))] == Decimal("4.02073")
    assert chave[(CAP, "Sul", date(2022, 12, 12))] == Decimal("3.45000")
    assert chave[(CAP, "Centro-Oeste", date(2021, 12, 13))] is None


# --- sobreposições -------------------------------------------------------------

def test_sobreposicao_dentro_do_arquivo():
    novos = [semana(date(2023, 1, 9), 1), semana(date(2023, 1, 12), 1)]
    (erro,) = importadores.sobreposicoes(novos, {})
    assert "do próprio arquivo" in erro and "12/01/2023" in erro


def test_sobreposicao_com_semana_cadastrada():
    existente = semana(date(2023, 1, 9), 1)
    existentes = {(CAP, existente["vigencia_inicio"], "Nordeste"): existente}
    (erro,) = importadores.sobreposicoes([semana(date(2023, 1, 12), 1)], existentes)
    assert "já cadastrada" in erro and "09/01/2023" in erro


def test_mesma_semana_do_arquivo_substitui_a_cadastrada():
    existente = semana(date(2023, 1, 9), 1, dias=10)
    existentes = {(CAP, existente["vigencia_inicio"], "Nordeste"): existente}
    novos = [semana(date(2023, 1, 9), 1), semana(date(2023, 1, 16), 1)]
    assert importadores.sobreposicoes(novos, existentes) == []


def test_sobreposicao_entre_cadastradas_nao_e_do_arquivo():
    """Só se reporta o que envolve o arquivo; o banco já proíbe o resto."""
    a, b = semana(date(2023, 1, 9), 1), semana(date(2023, 1, 12), 1)
    existentes = {(CAP, s["vigencia_inicio"], "Nordeste"): s for s in (a, b)}
    assert importadores.sobreposicoes([semana(date(2023, 2, 6), 1)], existentes) == []


# --- IGP-DI ----------------------------------------------------------------------

def _planilha(linhas, aba=importadores.ABA_IGP) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = aba
    for linha in linhas:
        ws.append(linha)
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def test_template_gerado_e_reimportavel():
    registros = [mes_igp(date(2023, 1, 1), "1143.861"), mes_igp(date(2022, 1, 1), "1110.398")]
    conteudo = importadores.gerar_template_igp_di(registros)
    leitura = importadores.ler_igp_di(conteudo)
    assert [(r["mes_ref"], r["valor"]) for r in leitura.registros] == [
        (date(2022, 1, 1), Decimal("1110.398")),
        (date(2023, 1, 1), Decimal("1143.861")),
    ]
    assert all(r["indice"] == INDICE_IGP_DI and r["base_label"] == BASE_IGP_DI for r in leitura.registros)
    wb = openpyxl.load_workbook(io.BytesIO(conteudo))
    assert wb.sheetnames == [importadores.ABA_IGP, importadores.ABA_INSTRUCOES]
    assert wb[importadores.ABA_IGP]["A2"].number_format == "mmm/yyyy"


def test_template_vazio_tem_so_o_cabecalho():
    wb = openpyxl.load_workbook(io.BytesIO(importadores.gerar_template_igp_di([])))
    assert [c.value for c in wb[importadores.ABA_IGP][1]] == ["Mês", "Valor"]
    assert wb[importadores.ABA_IGP].max_row == 1


def test_igp_aceita_mes_como_texto_e_ignora_linhas_em_branco():
    conteudo = _planilha([["Mês", "Valor"], ["03/2023", 1150.5], [None, None], [datetime(2023, 4, 1), 1151]])
    leitura = importadores.ler_igp_di(conteudo)
    assert [r["mes_ref"] for r in leitura.registros] == [date(2023, 3, 1), date(2023, 4, 1)]
    assert leitura.registros[1]["valor"] == Decimal("1151")


@pytest.mark.parametrize(
    "linhas,trecho",
    [
        ([["Mês", "Valor"], ["13/2023", 1]], "linha 2: mês"),
        ([["Mês", "Valor"], ["01/2023", 0]], "linha 2: valor"),
        ([["Mês", "Valor"], ["01/2023", "mil"]], "linha 2: valor"),
        ([["Mês", "Valor"], ["01/2023", 1], ["01/2023", 2]], "linha 3: o mês 01/2023 já aparece na linha 2"),
    ],
)
def test_igp_linha_invalida_recusa_o_arquivo(linhas, trecho):
    with pytest.raises(ArquivoInvalido) as erro:
        importadores.ler_igp_di(_planilha(linhas))
    assert any(trecho in e for e in erro.value.erros)


def test_igp_cabecalho_ou_aba_errados():
    with pytest.raises(ArquivoInvalido):
        importadores.ler_igp_di(_planilha([["Data", "Índice"], ["01/2023", 1]]))
    with pytest.raises(ArquivoInvalido):
        importadores.ler_igp_di(_planilha([["Mês", "Valor"]], aba="Plan1"))
    with pytest.raises(ArquivoInvalido):
        importadores.ler_igp_di(b"nao e xlsx")
    with pytest.raises(ArquivoInvalido):
        importadores.ler_igp_di(_planilha([["Mês", "Valor"]]))

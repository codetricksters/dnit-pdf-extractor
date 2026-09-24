"""API do cálculo: JSON e planilha saem da mesma função, com simulação de região."""

import io
from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from app.routers.api.calculo import cabecalho_avisos, nome_do_arquivo
from app.services import contratos_repo, indices_repo, medicoes_repo, template_repo, xlsx_drawings
from app.services.reequilibrio_layout import ABA

from .indices_factory import mes_igp, semana

HEADER = {
    "Contrato": "15 00716/2022 - HWN ENGENHARIA LTDA",
    "Data Base": "01/01/2022",
    "Número do Processo": "50615.000404/2022-67",
}
BASE = {"Nordeste": "4.02073", "Sul": "4.29019"}
JANEIRO = {"Nordeste": "3.28568", "Sul": "3.45"}


def _contrato(*, com_indices=True) -> int:
    contrato_id = contratos_repo.registrar_do_pdf(HEADER)
    medicoes_repo.gravar_itens(
        contrato_id,
        [{"Serviço": "8300980", "Descrição": "AQUISIÇÃO DE CAP 50/70",
          "Valor a PI Líquido": 208133.17, "Fator": -0.1839,
          "Período Líquido": "01/02/2023 - 28/02/2023", "Source_File": "1ª MP.pdf"}],
    )
    indices_repo.gravar_precos_anp(
        [semana(date(2023, 1, 9), JANEIRO["Nordeste"])]
    )
    contratos_repo.atualizar(contrato_id, {}, {"CAP": "Nordeste", "EMULSOES": "Nordeste"})
    if com_indices:
        for regiao in BASE:
            if regiao != "Nordeste":
                indices_repo.gravar_precos_anp([semana(date(2023, 1, 9), JANEIRO[regiao], regiao=regiao)])
            indices_repo.gravar_precos_anp([semana(date(2021, 12, 13), BASE[regiao], regiao=regiao)])
        indices_repo.gravar_indices_mensais(
            [mes_igp(date(2022, 1, 1), "1110.398"), mes_igp(date(2023, 2, 1), "1144.271")]
        )
    template_repo.garantir_semente()
    return contrato_id


def _delta(regiao: str) -> Decimal:
    return Decimal(JANEIRO[regiao]) / Decimal(BASE[regiao]) - 1


async def test_calculo_json(client):
    contrato_id = _contrato()
    resposta = await client.get(f"/api/v1/contratos/{contrato_id}/calculo")
    assert resposta.status_code == 200, resposta.text
    corpo = resposta.json()
    assert corpo["contrato"] == {"id": contrato_id, "numero": "15 00716/2022"}
    assert corpo["parametros"]["data_base"] == "2022-01-01"
    assert corpo["parametros"]["regioes"] == {"CAP": "Nordeste", "EMULSOES": "Nordeste"}
    assert corpo["parametros"]["simulacao"] is False
    assert corpo["parametros"]["lucro"] == "0.0511"

    (familia,) = corpo["familias"]
    assert familia["familia"] == "CAP" and familia["rotulo"] == "Aquisição de CAP"
    (linha,) = familia["produtos"][0]["linhas"]
    assert linha["mes"] == "2023-02-01"
    assert linha["a"] == "208133.17" and linha["b"] == "-38275.68"
    assert abs(Decimal(linha["d"]) - _delta("Nordeste")) < Decimal("1e-20")
    assert Decimal(corpo["total"]) == Decimal(linha["f"])


async def test_simulacao_nao_altera_o_cadastro(client):
    contrato_id = _contrato()
    corpo = (
        await client.get(f"/api/v1/contratos/{contrato_id}/calculo", params={"regiao_cap": "sul"})
    ).json()
    assert corpo["parametros"]["simulacao"] is True
    assert corpo["parametros"]["regioes"]["CAP"] == "Sul"
    linha = corpo["familias"][0]["produtos"][0]["linhas"][0]
    assert abs(Decimal(linha["d"]) - _delta("Sul")) < Decimal("1e-20")
    assert "Simulação: CAP calculado com a região Sul (cadastro: Nordeste)." in corpo["avisos"]

    cadastro = (await client.get(f"/api/v1/contratos/{contrato_id}")).json()
    assert cadastro["regioes"]["CAP"] == "Nordeste"


async def test_planilha_tem_formulas_vivas_e_a_equacao(client):
    contrato_id = _contrato()
    resposta = await client.get(f"/api/v1/contratos/{contrato_id}/planilha")
    assert resposta.status_code == 200, resposta.text
    assert 'filename="Reequilibrio_15-00716-2022.xlsx"' in resposta.headers["content-disposition"]

    ws = openpyxl.load_workbook(io.BytesIO(resposta.content))[ABA]
    linhas = [
        r for r in range(1, ws.max_row + 1)
        if str(ws.cell(r, 6).value or "").startswith("=TRUNC(")
    ]
    assert linhas
    r = linhas[0]
    assert ws.cell(r, 8).value == f"=D{r}*G{r}"
    assert ws.cell(r, 9).value == f"=H{r}-F{r}"
    assert ws.cell(r, 10).value == f"=I{r}*(1-0.0511)"
    assert isinstance(ws.cell(r, 7).value, float)
    assert xlsx_drawings.contem_equacao(resposta.content, ABA)


async def test_planilha_simulada_no_nome_e_nos_avisos(client):
    contrato_id = _contrato()
    resposta = await client.get(
        f"/api/v1/contratos/{contrato_id}/planilha", params={"regiao_cap": "Sul"}
    )
    assert resposta.status_code == 200
    assert "Reequilibrio_15-00716-2022_SIMULACAO_CAP-Sul.xlsx" in resposta.headers["content-disposition"]
    avisos = dict(resposta.headers.raw)[b"x-avisos"].decode("latin-1")
    assert "Simulação: CAP calculado com a região Sul (cadastro: Nordeste)." in avisos


async def test_contrato_inexistente_e_404(client):
    for rota in ("calculo", "planilha"):
        resposta = await client.get(f"/api/v1/contratos/999/{rota}")
        assert resposta.status_code == 404


async def test_indices_ausentes_sao_422_com_faltando(client):
    contrato_id = _contrato(com_indices=False)
    for rota in ("calculo", "planilha"):
        resposta = await client.get(f"/api/v1/contratos/{contrato_id}/{rota}")
        assert resposta.status_code == 422
        corpo = resposta.json()
        assert corpo["faltando"]
        assert all(f in corpo["detail"] for f in corpo["faltando"])


async def test_regiao_sem_precos_na_simulacao_e_422(client):
    contrato_id = _contrato()
    resposta = await client.get(
        f"/api/v1/contratos/{contrato_id}/calculo", params={"regiao_emulsoes": "Centro-Oeste"}
    )
    assert resposta.status_code == 422
    assert resposta.json()["faltando"] == ["regiao_emulsoes"]


@pytest.mark.parametrize(
    "simuladas,esperado",
    [
        ({}, "Reequilibrio_15-00716-2022.xlsx"),
        ({"CAP": "Sul"}, "Reequilibrio_15-00716-2022_SIMULACAO_CAP-Sul.xlsx"),
        (
            {"CAP": "Sul", "EMULSOES": "Centro-Oeste"},
            "Reequilibrio_15-00716-2022_SIMULACAO_CAP-Sul_EMULSOES-Centro-Oeste.xlsx",
        ),
    ],
)
def test_nome_do_arquivo(simuladas, esperado):
    assert nome_do_arquivo("15 00716/2022", simuladas) == esperado


def test_cabecalho_avisos_e_latin1():
    texto = cabecalho_avisos(["Região Sul", "ΔP – sem índice…"])
    texto.encode("latin-1")
    assert texto.startswith("Região Sul | ")

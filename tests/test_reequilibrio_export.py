"""The exported spreadsheet: live formulas, surviving equation, right totals."""

import io
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from app.services import contratos_repo, medicoes_repo, reequilibrio_export, xlsx_drawings
from app.services.delta_p import FAMILIA_CAP, FAMILIA_EMULSOES
from app.services.reequilibrio_export import (
    ExportacaoImpossivel,
    Grupo,
    Linha,
    gerar_planilha,
    montar_grupos,
)
from app.services.reequilibrio_layout import ABA, PRIMEIRA_LINHA, ROTULO_SUBTOTAL, ROTULO_TOTAL

TEMPLATE = Path("app/templates_xlsx/reequilibrio_template.xlsx")

HEADER = {
    "Contrato": "15 00716/2022 - HWN ENGENHARIA LTDA Índices I0 I1 K",
    "Data Base": "01/01/2022",
    "Período Líquido": "01/03/2023 - 31/03/2023",
    "Número do Processo": "50615.000404/2022-67",
}

CONTRATO = {
    "numero": "15 00716/2022",
    "edital": "000433/2022-15",
    "rodovia": "BR-316/MA",
    "trecho": "DIV PA/MA - ENTR MA-306/206",
    "subtrecho": "NOVA OLINDA DO MARANHÃO",
    "segmento": "KM 0,00 A KM 188,70",
    "extensao": Decimal("188.7"),
    "contratada": "HWN Engenharia LTDA",
    "numero_processo": "50615.000404/2022-67",
    "data_base": date(2022, 1, 1),
}


@pytest.fixture
def template() -> bytes:
    return TEMPLATE.read_bytes()


def _grupos():
    return [
        Grupo(
            descricao="AQUISIÇÃO DE CAP 50/70",
            familia=FAMILIA_CAP,
            linhas=[
                Linha(date(2023, 1, 1), Decimal("208133.17"), Decimal("-0.1839"), Decimal("-0.0758")),
                Linha(date(2023, 2, 1), Decimal("100000.00"), Decimal("-0.1839"), Decimal("-0.1828")),
            ],
        ),
        Grupo(
            descricao="AQUISIÇÃO DE EMULSÃO ASFÁLTICA RR-1C",
            familia=FAMILIA_EMULSOES,
            linhas=[
                Linha(date(2023, 1, 1), Decimal("25130.23"), Decimal("-0.1369"), Decimal("-0.0493")),
            ],
        ),
    ]


def _aba(conteudo: bytes):
    return openpyxl.load_workbook(io.BytesIO(conteudo))[ABA]


async def test_template_versionado_tem_o_layout_esperado(template):
    ws = _aba(template)
    assert ws["B3"].value == "Contrato: "
    assert ws["B14"].value == "MEDIÇÃO (MÊS)"
    assert ws["J16"].value == "f = e * (1-(5,11/100))"
    assert xlsx_drawings.contem_equacao(template, ABA) is True


async def test_colunas_calculadas_saem_como_formulas(template):
    """The whole point of the export: the user audits by clicking the cell."""
    ws = _aba(gerar_planilha(CONTRATO, _grupos(), template))
    linha = PRIMEIRA_LINHA + 1  # first data row, after the group header

    assert ws.cell(linha, 6).value == f"=TRUNC(E{linha}*D{linha},2)"
    assert ws.cell(linha, 8).value == f"=D{linha}*G{linha}"
    assert ws.cell(linha, 9).value == f"=H{linha}-F{linha}"
    assert ws.cell(linha, 10).value == f"=I{linha}*(1-0.0511)"


async def test_delta_p_e_o_unico_valor_literal(template):
    """ΔP is computed by the app; every other derived column is a formula.

    Read back as floats because that is how openpyxl parses a number cell; the
    exactness of what was *written* is the next test's subject.
    """
    ws = _aba(gerar_planilha(CONTRATO, _grupos(), template))
    linha = PRIMEIRA_LINHA + 1

    assert ws.cell(linha, 7).value == pytest.approx(-0.0758)
    assert ws.cell(linha, 4).value == pytest.approx(208133.17)  # do PDF
    assert ws.cell(linha, 5).value == pytest.approx(-0.1839)  # do PDF


async def test_valores_chegam_a_planilha_sem_perda(template):
    """The Decimal carried through the DB and ΔP reaches the cell unchanged.

    Compared as float64 because that is all an xlsx number cell can hold; the
    exactness guarantee belongs to the database and the arithmetic, not to the
    spreadsheet's storage.
    """
    grupos = [
        Grupo(
            descricao="AQUISIÇÃO DE CAP 50/70",
            familia=FAMILIA_CAP,
            linhas=[
                Linha(
                    date(2023, 1, 1),
                    Decimal("1872240.49"),
                    Decimal("-0.1839"),
                    Decimal("-0.0758"),
                )
            ],
        )
    ]
    ws = _aba(gerar_planilha(CONTRATO, grupos, template))
    assert ws.cell(18, 4).value == float(Decimal("1872240.49"))
    assert ws.cell(18, 7).value == float(Decimal("-0.0758"))


async def test_memoria_de_calculo_sobrevive_a_exportacao(template):
    """Guards the regression openpyxl causes on its own: text boxes vanish."""
    conteudo = gerar_planilha(CONTRATO, _grupos(), template)

    formas = xlsx_drawings.extrair_formas(conteudo, ABA)
    assert len(formas) == 1
    assert len(re.findall(r"<a:t>", formas.blocos[0])) >= 10
    assert xlsx_drawings.contem_equacao(conteudo, ABA) is True
    # The logo, which openpyxl does preserve, is not duplicated by reinjection.
    assert len(_aba(conteudo)._images) == 1


async def test_cabecalho_do_contrato_e_preenchido(template):
    ws = _aba(gerar_planilha(CONTRATO, _grupos(), template))
    assert ws["C3"].value == "15 00716/2022"
    assert ws["C5"].value == "BR-316/MA"
    assert ws["C9"].value == pytest.approx(188.7)
    assert ws["C12"].value == datetime(2022, 1, 1)  # openpyxl lê datas como datetime
    # Período stays a template formula over the months actually written.
    assert str(ws["C13"].value).startswith("=PROPER(")


async def test_agrupamento_subtotais_e_total(template):
    ws = _aba(gerar_planilha(CONTRATO, _grupos(), template))

    assert ws.cell(17, 2).value == "AQUISIÇÃO DE CAP 50/70"
    assert ws.cell(18, 3).value == "=B17"  # descrição aponta para o grupo
    assert ws.cell(20, 8).value == ROTULO_SUBTOTAL
    assert ws.cell(20, 10).value == "=SUM(J18:J19)"

    assert ws.cell(21, 2).value == "AQUISIÇÃO DE EMULSÃO ASFÁLTICA RR-1C"
    assert ws.cell(23, 10).value == "=SUM(J22:J22)"

    assert ws.cell(24, 8).value == ROTULO_TOTAL
    assert ws.cell(24, 10).value == "=J20+J23"


async def test_formatacao_das_linhas_modelo_e_reaproveitada(template):
    """Style comes from the template so the user can restyle without code."""
    modelo = _aba(template)
    ws = _aba(gerar_planilha(CONTRATO, _grupos(), template))

    assert ws.cell(18, 4).number_format == modelo.cell(18, 4).number_format
    assert ws.cell(18, 2).number_format == modelo.cell(18, 2).number_format
    assert ws.cell(18, 4).border.left.style == modelo.cell(18, 4).border.left.style
    assert ws.cell(17, 2).font.b == modelo.cell(17, 2).font.b


async def test_linhas_modelo_nao_vazam_para_a_planilha_final(template):
    ws = _aba(gerar_planilha(CONTRATO, [], template))
    assert ws.max_row == PRIMEIRA_LINHA  # only the total row
    assert ws.cell(PRIMEIRA_LINHA, 8).value == ROTULO_TOTAL


# --- montar_grupos ---------------------------------------------------------


def _item(descricao, familia, codigo, mes, valor, fator="-0.1839"):
    return {
        "descricao_export": descricao,
        "familia": familia,
        "codigo_servico": codigo,
        "mes_medicao": mes,
        "valor_pi": Decimal(valor),
        "fator": Decimal(fator),
    }


async def test_codigos_diferentes_do_mesmo_produto_somam_no_mes():
    """One row per month per product, as in the reference layout."""
    deltas = {(FAMILIA_CAP, date(2023, 1, 1)): Decimal("-0.0758")}
    grupos, descartadas = montar_grupos(
        [
            _item("CAP", FAMILIA_CAP, "60112", date(2023, 1, 1), "100.00"),
            _item("CAP", FAMILIA_CAP, "92704", date(2023, 1, 1), "50.50"),
        ],
        deltas,
    )
    assert descartadas == 0
    assert len(grupos) == 1
    assert [l.valor_pi for l in grupos[0].linhas] == [Decimal("150.50")]


async def test_fatores_diferentes_nao_sao_misturados():
    """Blending different fatores would corrupt column F."""
    deltas = {(FAMILIA_CAP, date(2023, 1, 1)): Decimal("-0.0758")}
    grupos, _ = montar_grupos(
        [
            _item("CAP", FAMILIA_CAP, "60112", date(2023, 1, 1), "100.00", "-0.1839"),
            _item("CAP", FAMILIA_CAP, "92704", date(2023, 1, 1), "50.00", "-0.2252"),
        ],
        deltas,
    )
    assert len(grupos[0].linhas) == 2
    assert {l.fator for l in grupos[0].linhas} == {Decimal("-0.1839"), Decimal("-0.2252")}


async def test_mesma_medicao_em_dois_arquivos_nao_dobra_o_total():
    """A PDF re-uploaded under another name must not double the value."""
    deltas = {(FAMILIA_CAP, date(2023, 1, 1)): Decimal("-0.0758")}
    item = _item("CAP", FAMILIA_CAP, "60112", date(2023, 1, 1), "100.00")
    grupos, descartadas = montar_grupos([item, dict(item)], deltas)
    assert descartadas == 1
    assert [l.valor_pi for l in grupos[0].linhas] == [Decimal("100.00")]


# --- exportar (ponta a ponta com o banco) ----------------------------------


async def _contrato_com_medicoes():
    contrato_id = contratos_repo.registrar_do_pdf(HEADER)
    contratos_repo.definir_regiao("15 00716/2022", FAMILIA_CAP, "Nordeste")
    contratos_repo.definir_regiao("15 00716/2022", FAMILIA_EMULSOES, "Nordeste")
    medicoes_repo.gravar_itens(
        contrato_id,
        [
            {
                "Serviço": "8300980",
                "Descrição": "AQUISIÇÃO DE CIMENTO ASFÁLTICO CAP 50/70",
                "Valor a PI Líquido": 208133.17,
                "Fator": -0.1839,
                "Período Líquido": "01/02/2023 - 28/02/2023",
                "Source_File": "1ª MP.pdf",
            }
        ],
    )
    return contrato_id


async def test_exportar_sem_indices_explica_o_que_falta(template):
    await _contrato_com_medicoes()
    with pytest.raises(ExportacaoImpossivel) as erro:
        reequilibrio_export.exportar("15 00716/2022", template)
    assert "ΔP" in str(erro.value)


async def test_exportar_sem_regiao_pede_a_regiao(template):
    contrato_id = contratos_repo.registrar_do_pdf(HEADER)
    medicoes_repo.gravar_itens(
        contrato_id,
        [
            {
                "Serviço": "8300980",
                "Descrição": "CAP",
                "Valor a PI Líquido": 10.0,
                "Fator": -0.1,
                "Período Líquido": "01/02/2023 - 28/02/2023",
                "Source_File": "1ª MP.pdf",
            }
        ],
    )
    with pytest.raises(ExportacaoImpossivel) as erro:
        reequilibrio_export.exportar("15 00716/2022", template)
    assert "região" in str(erro.value).lower()


async def test_indice_faltante_e_relatado_uma_vez_por_motivo():
    """One missing value is one thing to fix, not one per measurement.

    The ΔP base month is shared by every line, so an absent base índice used to
    produce an identical bullet for each of the 52 items.
    """
    contrato_id = await _contrato_com_medicoes()
    medicoes_repo.gravar_itens(
        contrato_id,
        [
            {
                "Serviço": "8300980",
                "Descrição": "AQUISIÇÃO DE CIMENTO ASFÁLTICO CAP 50/70",
                "Valor a PI Líquido": 10.0,
                "Fator": -0.1839,
                "Período Líquido": f"01/0{m}/2023 - 28/0{m}/2023",
                "Source_File": f"{m}ª MP.pdf",
            }
            for m in range(3, 8)
        ],
    )
    contrato = contratos_repo.buscar("15 00716/2022")
    itens = medicoes_repo.itens_para_export(contrato_id)

    _, faltando = reequilibrio_export.calcular_deltas(contrato, itens)

    assert len(itens) > len(faltando)
    assert len(faltando) == len(set(faltando))


async def test_exportar_contrato_inexistente(template):
    with pytest.raises(ExportacaoImpossivel):
        reequilibrio_export.exportar("99 99999/9999", template)


async def test_exportar_sem_codigos_associados(template):
    contratos_repo.registrar_do_pdf(HEADER)
    with pytest.raises(ExportacaoImpossivel) as erro:
        reequilibrio_export.exportar("15 00716/2022", template)
    assert "associado" in str(erro.value)

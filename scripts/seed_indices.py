"""Initial load of the price indices.

One-off, run from the command line — never from a request path. From then on the
users keep the tables up to date through the interface as new figures are
published.

    uv run python scripts/seed_indices.py

Two sources, used only as a *model* for the data the application stores:

* ``data/precos-medios-ponderados-semanais-2013.xls`` — ANP weekly producer
  prices. Legacy BIFF, so it needs ``xlrd`` (a dev dependency; the application
  itself never reads this file). Dates come as Excel serials and missing
  quotations as ``'***'``.
* ``data/Reequilíbrio - 26 - Contrato 716-22.xlsx``, sheet ``IGP - DI`` — the
  monthly DNIT indices, in wide format (one column per month), transposed to the
  long format the database uses.
"""

import argparse
import sys
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app import config, db  # noqa: E402
from app.services.delta_p import ANP_PRODUTO_CAP  # noqa: E402
from app.services.indices_repo import (  # noqa: E402
    gravar_indices_mensais,
    gravar_precos_anp,
)

ARQUIVO_ANP = config.PROJECT_ROOT / "data" / "precos-medios-ponderados-semanais-2013.xls"
ARQUIVO_DNIT = config.PROJECT_ROOT / "data" / "Reequilíbrio - 26 - Contrato 716-22.xlsx"

# Column order of the ANP sheet, from row 8 of its two-row header.
REGIOES_ANP = ["Norte", "Nordeste", "Centro-Oeste", "Sul", "Sudeste", "Brasil"]
PRIMEIRA_LINHA_ANP = 9

# Excel's 1900 date system, with the usual two-day offset for its leap-year bug.
EPOCA_EXCEL = date(1899, 12, 30)

CENTAVOS_ANP = Decimal("0.00001")  # the ANP publishes 5 decimal places


def _data_de_serial(serial: float) -> date:
    return EPOCA_EXCEL + timedelta(days=int(serial))


def _preco(valor) -> Decimal | None:
    """Convert a cell to Decimal, or None when there is no quotation.

    ``'***'`` means the ANP published no price for that region that week — a
    real absence, stored as NULL rather than zero. The quantize step removes the
    binary float noise xlrd hands back (1.2935999999999999), so the stored value
    matches the published figure exactly.
    """
    if isinstance(valor, str) or valor in ("", None):
        return None
    return Decimal(str(valor)).quantize(CENTAVOS_ANP)


def carregar_anp(produtos: set[str] | None) -> int:
    import xlrd

    if not ARQUIVO_ANP.exists():
        print(f"[ANP] arquivo não encontrado: {ARQUIVO_ANP}")
        return 0

    sheet = xlrd.open_workbook(str(ARQUIVO_ANP)).sheet_by_index(0)
    registros: list[dict] = []
    for i in range(PRIMEIRA_LINHA_ANP, sheet.nrows):
        linha = sheet.row_values(i)
        produto = str(linha[0]).strip()
        if not produto or produtos is not None and produto not in produtos:
            continue
        if not isinstance(linha[1], float) or not isinstance(linha[2], float):
            continue
        inicio, fim = _data_de_serial(linha[1]), _data_de_serial(linha[2])
        for deslocamento, regiao in enumerate(REGIOES_ANP):
            registros.append(
                {
                    "produto": produto,
                    "vigencia_inicio": inicio,
                    "vigencia_fim": fim,
                    "regiao": regiao,
                    "preco": _preco(linha[3 + deslocamento]),
                }
            )

    gravar_precos_anp(registros)
    print(f"[ANP] {len(registros)} registros gravados")
    return len(registros)


def carregar_indices_dnit() -> int:
    import openpyxl

    if not ARQUIVO_DNIT.exists():
        print(f"[DNIT] arquivo não encontrado: {ARQUIVO_DNIT}")
        return 0

    ws = openpyxl.load_workbook(ARQUIVO_DNIT, data_only=True)["IGP - DI"]

    # Row 3 holds one month per column, from column C onwards.
    meses: dict[int, date] = {}
    for celula in ws[3]:
        if isinstance(celula.value, datetime):
            meses[celula.column] = celula.value.date().replace(day=1)

    registros: list[dict] = []
    for linha in ws.iter_rows(min_row=4):
        nome = linha[0].value
        if not isinstance(nome, str) or not nome.strip():
            continue
        base_label = linha[1].value if isinstance(linha[1].value, str) else None
        for celula in linha:
            mes = meses.get(celula.column)
            if mes is None or not isinstance(celula.value, (int, float)):
                continue
            registros.append(
                {
                    "indice": nome.strip(),
                    "base_label": base_label,
                    "mes_ref": mes,
                    "valor": Decimal(str(celula.value)),
                }
            )

    gravar_indices_mensais(registros)
    indices = sorted({r["indice"] for r in registros})
    print(f"[DNIT] {len(registros)} registros gravados, {len(indices)} índices")

    if registros:
        primeiro = min(r["mes_ref"] for r in registros)
        print(
            f"[DNIT] atenção: a série começa em {primeiro:%m/%Y}. O ΔP das "
            "emulsões precisa do IGP-DI do mês da Data Base do contrato; para "
            "contratos com data base anterior a esse mês, informe o valor pela "
            "interface antes de calcular."
        )
    return len(registros)


async def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--todos-produtos",
        action="store_true",
        help="carrega os 29 produtos da ANP; por padrão só o CAP 50/70, "
        "único usado no cálculo do ΔP",
    )
    args = parser.parse_args()

    await db.open_pools()
    try:
        await db.apply_migrations()
        carregar_anp(None if args.todos_produtos else {ANP_PRODUTO_CAP})
        carregar_indices_dnit()
    finally:
        await db.close_pools()


if __name__ == "__main__":
    import asyncio

    asyncio.run(main())

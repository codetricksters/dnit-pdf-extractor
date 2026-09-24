"""Ponta a ponta, só pela API, com índices reais e itens reais de medição.

O oráculo é a aba ``CÁLCULO DA VARIAÇÃO DE PREÇOS`` da planilha de referência
(contrato 716-22, Data Base jan/2022, Nordeste), calculada pelo Excel sem
nenhum código desta aplicação. Se o ΔP divergir, é divergência de regra a levar
ao usuário: o oráculo não é ajustado para o teste passar.
"""

import csv
import io
import json
from decimal import ROUND_DOWN, Decimal
from pathlib import Path

import openpyxl
import pytest

from app.services import template_repo, xlsx_drawings
from app.services.file_processor import _persistir
from app.services.reequilibrio_layout import ABA

FIXTURES = Path(__file__).parent / "fixtures"
NUMERO = "99 99999/2099"
TOLERANCIA = Decimal("1e-9")
LUCRO = Decimal("0.0511")
CENTAVO = Decimal("0.01")


def _oraculo() -> dict[str, dict[str, Decimal]]:
    with (FIXTURES / "delta_p_referencia.csv").open(encoding="utf-8") as f:
        return {
            linha["mes"]: {"CAP": Decimal(linha["delta_cap"]), "EMULSOES": Decimal(linha["delta_emul"])}
            for linha in csv.DictReader(f)
        }


async def _importar(client, rota: str, arquivo: Path) -> dict:
    resposta = await client.post(
        f"/api/v1/indices/{rota}/importar",
        files={"arquivo": (arquivo.name, arquivo.read_bytes())},
    )
    assert resposta.status_code == 200, resposta.text
    return resposta.json()


async def _preparar(client) -> int:
    # 1. Índices reais, pelos endpoints de importação.
    anp = await _importar(client, "anp", FIXTURES / "anp_semanal.xls")
    assert anp["inseridos"] == 60114
    igp = await _importar(client, "igp-di", FIXTURES / "igp_di.xlsx")
    assert igp["periodo"]["de"] == "2022-01-01"

    # 2. O contrato, pelo mesmo caminho de um PDF processado.
    ficticio = json.loads((FIXTURES / "contrato_ficticio.json").read_text(encoding="utf-8"))
    for medicao in ficticio["arquivos"]:
        _persistir({"header": ficticio["header"], "rows": medicao["rows"]}, medicao["arquivo"], "e2e")
    (resumo,) = (await client.get("/api/v1/contratos", params={"numero": NUMERO})).json()
    contrato_id = resumo["id"]

    # 3. Cadastro completo, com a região em grafia diferente da do arquivo ANP.
    resposta = await client.patch(
        f"/api/v1/contratos/{contrato_id}",
        json={
            "edital": "999/2099-99", "rodovia": "BR-999", "trecho": "Trecho fictício",
            "subtrecho": "Subtrecho fictício", "segmento": "km 0,0 ao km 10,0",
            "extensao": "10.0", "contratada": "CONSTRUTORA FICTÍCIA LTDA",
            "regioes": {"CAP": "nordeste", "EMULSOES": "Nordeste"},
        },
    )
    assert resposta.status_code == 200, resposta.text
    assert resposta.json()["regioes"] == {"CAP": "Nordeste", "EMULSOES": "Nordeste"}

    # 4. Catálogo do zero: a migração traz associações prontas; o teste as remove
    #    para provar que só o que o usuário associa entra no cálculo.
    for codigo in (await client.get("/api/v1/codigos", params={"associado": "true"})).json():
        assert (await client.delete(f"/api/v1/codigos/{codigo['codigo']}")).status_code == 204

    # 5. Dois produtos customizados e só dois códigos associados.
    cap = await client.post(
        "/api/v1/produtos", json={"descricao_export": "Aquisição de CAP 50/70", "familia": "CAP"}
    )
    emulsao = await client.post(
        "/api/v1/produtos",
        json={"descricao_export": "Aquisição de Emulsão RR-1C", "familia": "EMULSOES"},
    )
    assert cap.status_code == emulsao.status_code == 201
    for codigo, produto in (("60112", cap), ("29083", emulsao)):
        resposta = await client.put(
            f"/api/v1/codigos/{codigo}", json={"produto_id": produto.json()["id"]}
        )
        assert resposta.status_code == 200, resposta.text

    template_repo.garantir_semente()
    return contrato_id


def _linhas(corpo: dict):
    for familia in corpo["familias"]:
        for produto in familia["produtos"]:
            for linha in produto["linhas"]:
                yield familia["familia"], produto["descricao"], linha


async def test_ponta_a_ponta(client):
    contrato_id = await _preparar(client)
    oraculo = _oraculo()

    # --- cálculo em JSON ---------------------------------------------------
    resposta = await client.get(f"/api/v1/contratos/{contrato_id}/calculo")
    assert resposta.status_code == 200, resposta.text
    corpo = resposta.json()
    assert corpo["parametros"]["data_base"] == "2022-01-01"
    assert corpo["parametros"]["simulacao"] is False

    # Só os dois produtos associados; os outros códigos ficam fora sem bloquear.
    assert {(f["familia"], p["descricao"]) for f in corpo["familias"] for p in f["produtos"]} == {
        ("CAP", "Aquisição de CAP 50/70"),
        ("EMULSOES", "Aquisição de Emulsão RR-1C"),
    }

    linhas = list(_linhas(corpo))
    meses = {linha["mes"] for _, _, linha in linhas}
    assert meses <= {f"2023-{m:02d}-01" for m in range(1, 13)}
    assert "2023-01-01" in meses

    # reequilibrio_export.serializar() soma f por produto (o subtotal do
    # produto), acumula esse subtotal na família e no total ao mesmo tempo —
    # o total não é a soma dos subtotais de família, é a soma direta dos
    # subtotais de produto. A comparação replica essa mesma ordem para ser bit
    # exata em cada nível, e ainda confere o subtotal de família separadamente.
    total = Decimal(0)
    for familia in corpo["familias"]:
        subtotal_familia = Decimal(0)
        for produto in familia["produtos"]:
            subtotal_produto = Decimal(0)
            for linha in produto["linhas"]:
                a, fator, d = Decimal(linha["a"]), Decimal(linha["fator"]), Decimal(linha["d"])
                # ΔP contra o oráculo do Excel.
                assert abs(d - oraculo[linha["mes"]][familia["familia"]]) < TOLERANCIA, (
                    familia["familia"], linha["mes"]
                )
                # As fórmulas do art. 16, recalculadas aqui.
                b = (a * fator).quantize(CENTAVO, rounding=ROUND_DOWN)
                c = a * d
                e = c - b
                f = e * (1 - LUCRO)
                assert Decimal(linha["b"]) == b
                assert Decimal(linha["c"]) == c
                assert Decimal(linha["e"]) == e
                assert Decimal(linha["f"]) == f
                subtotal_produto += f
            assert Decimal(produto["subtotal"]) == subtotal_produto
            subtotal_familia += subtotal_produto
            total += subtotal_produto
        assert Decimal(familia["subtotal"]) == subtotal_familia
    assert Decimal(corpo["total"]) == total

    janeiro_cap = [l for fam, _, l in linhas if fam == "CAP" and l["mes"] == "2023-01-01"]
    assert [l["a"] for l in janeiro_cap] == ["21862.28"]
    # De março em diante cada PDF traz o 60112 duas vezes (uma linha zerada e a
    # linha com valor); o banco guarda a última. Ver a nota do Step 4.
    marco_cap = [l for fam, _, l in linhas if fam == "CAP" and l["mes"] == "2023-03-01"]
    assert [l["a"] for l in marco_cap] == ["13926.57"]

    # --- a planilha ----------------------------------------------------------
    resposta = await client.get(f"/api/v1/contratos/{contrato_id}/planilha")
    assert resposta.status_code == 200, resposta.text
    assert 'filename="Reequilibrio_99-99999-2099.xlsx"' in resposta.headers["content-disposition"]
    ws = openpyxl.load_workbook(io.BytesIO(resposta.content))[ABA]
    formulas = [
        r for r in range(1, ws.max_row + 1)
        if str(ws.cell(r, 6).value or "").startswith("=TRUNC(")
    ]
    assert len(formulas) == len(linhas)
    deltas = sorted(float(linha["d"]) for _, _, linha in linhas)
    assert sorted(ws.cell(r, 7).value for r in formulas) == pytest.approx(deltas, abs=1e-12)
    for r in formulas:
        assert ws.cell(r, 6).value == f"=TRUNC(E{r}*D{r},2)"
        assert ws.cell(r, 8).value == f"=D{r}*G{r}"
        assert ws.cell(r, 9).value == f"=H{r}-F{r}"
        assert ws.cell(r, 10).value == f"=I{r}*(1-0.0511)"
    assert xlsx_drawings.contem_equacao(resposta.content, ABA)

    # --- simulação de região -------------------------------------------------
    simulada = await client.get(
        f"/api/v1/contratos/{contrato_id}/planilha", params={"regiao_cap": "Sul"}
    )
    assert simulada.status_code == 200, simulada.text
    assert "SIMULACAO_CAP-Sul" in simulada.headers["content-disposition"]
    corpo_sul = (
        await client.get(f"/api/v1/contratos/{contrato_id}/calculo", params={"regiao_cap": "Sul"})
    ).json()
    assert corpo_sul["parametros"]["simulacao"] is True
    cap_nordeste = {l["mes"]: l["d"] for fam, _, l in linhas if fam == "CAP"}
    cap_sul = {l["mes"]: l["d"] for fam, _, l in _linhas(corpo_sul) if fam == "CAP"}
    assert cap_sul.keys() == cap_nordeste.keys()
    assert all(cap_sul[m] != cap_nordeste[m] for m in cap_sul)
    cadastro = (await client.get(f"/api/v1/contratos/{contrato_id}")).json()
    assert cadastro["regioes"] == {"CAP": "Nordeste", "EMULSOES": "Nordeste"}

    # --- recusas -------------------------------------------------------------
    # Centro-Oeste não tem cotação de CAP até maio/2023: o cálculo é recusado.
    sem_cotacao = await client.get(
        f"/api/v1/contratos/{contrato_id}/calculo", params={"regiao_cap": "Centro-Oeste"}
    )
    assert sem_cotacao.status_code == 422
    assert sem_cotacao.json()["faltando"]
    inexistente = await client.get(
        f"/api/v1/contratos/{contrato_id}/calculo", params={"regiao_cap": "Atlântida"}
    )
    assert inexistente.status_code == 422

"""Export templates, kept in the database.

Stored as ``BYTEA`` rather than as files in a volume so that the database backup
covers them: one artefact to save and restore instead of two. The user runs in
Docker and has no comfortable access to container files, so upload, download,
activation and deletion all happen through the interface.

A malformed template would produce a corrupt or empty spreadsheet without any
error, so uploads are validated and refused with a reason.
"""

import hashlib
import io
from pathlib import Path

import openpyxl

from ..db import acquire_sync
from .reequilibrio_layout import (
    ABA,
    CABECALHOS,
    LINHA_CABECALHO_FIM,
    LINHA_CABECALHO_INICIO,
    LINHA_GRUPO,
    LINHA_HEADER,
    LINHA_TOTAL,
)
from .xlsx_drawings import TemplateInvalido, contem_equacao

SEMENTE = Path(__file__).resolve().parents[1] / "templates_xlsx" / "reequilibrio_template.xlsx"


def _normalizar(texto) -> str:
    """Compare labels ignoring case and whitespace.

    PDF-style labels and Excel headers differ in trailing spaces and line
    breaks (``"VALOR\\n A PI"``), and refusing a template over a space would be
    obstructive.
    """
    return " ".join(str(texto or "").split()).casefold()


def validar(conteudo: bytes) -> None:
    """Raise ``TemplateInvalido`` if the export could not use this file."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(conteudo))
    except Exception as e:
        raise TemplateInvalido(
            "Não foi possível abrir o arquivo como planilha .xlsx."
        ) from e

    if ABA not in wb.sheetnames:
        raise TemplateInvalido(
            f"A planilha precisa ter uma aba chamada '{ABA}'. "
            f"Abas encontradas: {', '.join(wb.sheetnames)}."
        )
    ws = wb[ABA]

    rotulos = {
        linha: _normalizar(ws.cell(linha, 2).value)
        for linha in range(LINHA_CABECALHO_INICIO, LINHA_CABECALHO_FIM + 1)
    }
    if "contrato" not in rotulos[LINHA_CABECALHO_INICIO]:
        raise TemplateInvalido(
            f"A célula B{LINHA_CABECALHO_INICIO} deveria rotular o contrato; "
            "o cabeçalho do contrato precisa começar em B3."
        )
    if "data base" not in rotulos[12]:
        raise TemplateInvalido(
            "A célula B12 deveria rotular a Data base. O bloco B3:C13 está "
            "deslocado, e a exportação escreveria os valores nas linhas erradas."
        )

    for coluna, esperado in CABECALHOS.items():
        encontrado = _normalizar(ws.cell(LINHA_HEADER, coluna).value)
        if _normalizar(esperado) != encontrado:
            letra = ws.cell(LINHA_HEADER, coluna).column_letter
            raise TemplateInvalido(
                f"O cabeçalho da tabela está diferente em {letra}{LINHA_HEADER}: "
                f"esperado '{' '.join(esperado.split())}', encontrado "
                f"'{ws.cell(LINHA_HEADER, coluna).value}'."
            )

    # The model rows (group header, data, subtotal, total) are what the export
    # copies formatting from. The group header's horizontal merge is their
    # marker: a template where they were deleted in Excel loses it.
    faixas = {str(m) for m in ws.merged_cells.ranges}
    if ws.max_row < LINHA_TOTAL or f"B{LINHA_GRUPO}:J{LINHA_GRUPO}" not in faixas:
        raise TemplateInvalido(
            f"O template precisa conter as linhas-modelo {LINHA_GRUPO} a "
            f"{LINHA_TOTAL} (grupo, dados, subtotal e total), com B{LINHA_GRUPO}:"
            f"J{LINHA_GRUPO} mesclado. É de onde a exportação copia a formatação."
        )

    if not contem_equacao(conteudo, ABA):
        raise TemplateInvalido(
            "A memória de cálculo não foi encontrada. O template precisa conter "
            "a equação numa caixa de texto."
        )


def garantir_semente() -> None:
    """Install the packaged template the first time the app starts.

    From then on the database is the source of truth: the file on disk is not
    re-read, so an activation or an edit made through the interface is not
    undone by a restart.
    """
    with acquire_sync() as conn:
        existe = conn.execute("SELECT 1 FROM template LIMIT 1").fetchone()
        if existe:
            return
    salvar(SEMENTE.name, SEMENTE.read_bytes(), observacao="Template inicial", ativar=True)


def salvar(nome: str, conteudo: bytes, observacao: str | None = None, ativar: bool = True) -> int:
    validar(conteudo)
    sha = hashlib.sha256(conteudo).hexdigest()
    with acquire_sync() as conn:
        with conn.transaction():
            if ativar:
                conn.execute("UPDATE template SET ativo = false WHERE ativo")
            linha = conn.execute(
                "INSERT INTO template (nome, arquivo, tamanho, sha256, ativo, observacao) "
                "VALUES (%s, %s, %s, %s, %s, %s) RETURNING id",
                (nome, conteudo, len(conteudo), sha, ativar, observacao),
            ).fetchone()
    return linha["id"]


def listar() -> list[dict]:
    """Templates without their bytes, most recent first."""
    with acquire_sync() as conn:
        cur = conn.execute(
            "SELECT id, nome, tamanho, sha256, ativo, criado_em, observacao "
            "FROM template ORDER BY criado_em DESC, id DESC"
        )
        return [dict(r) for r in cur.fetchall()]


def buscar(template_id: int) -> dict | None:
    with acquire_sync() as conn:
        linha = conn.execute(
            "SELECT id, nome, arquivo, tamanho, sha256, ativo, criado_em, observacao "
            "FROM template WHERE id = %s",
            (template_id,),
        ).fetchone()
        return dict(linha) if linha else None


def ativo() -> dict | None:
    with acquire_sync() as conn:
        linha = conn.execute(
            "SELECT id, nome, arquivo, tamanho, sha256, criado_em FROM template "
            "WHERE ativo"
        ).fetchone()
        return dict(linha) if linha else None


def conteudo_ativo() -> bytes:
    atual = ativo()
    if atual is None:
        raise TemplateInvalido(
            "Nenhum template está ativo. Ative um template antes de exportar."
        )
    return bytes(atual["arquivo"])


def ativar(template_id: int) -> bool:
    with acquire_sync() as conn:
        with conn.transaction():
            # Clearing first keeps the partial unique index on (ativo) satisfied
            # at every point of the transaction.
            conn.execute("UPDATE template SET ativo = false WHERE ativo")
            cur = conn.execute(
                "UPDATE template SET ativo = true WHERE id = %s", (template_id,)
            )
            return cur.rowcount == 1


def excluir(template_id: int) -> None:
    """Remove a template.

    The active one cannot be removed, and one always remains — otherwise the
    export would be left without a usable base. Spreadsheets already generated
    are independent files and are unaffected.
    """
    with acquire_sync() as conn:
        with conn.transaction():
            linha = conn.execute(
                "SELECT ativo FROM template WHERE id = %s", (template_id,)
            ).fetchone()
            if linha is None:
                raise TemplateInvalido("Template não encontrado.")
            if linha["ativo"]:
                raise TemplateInvalido(
                    "Este é o template ativo. Ative outro antes de excluí-lo."
                )
            total = conn.execute("SELECT count(*) AS n FROM template").fetchone()["n"]
            if total <= 1:
                raise TemplateInvalido(
                    "É preciso manter ao menos um template cadastrado."
                )
            conn.execute("DELETE FROM template WHERE id = %s", (template_id,))

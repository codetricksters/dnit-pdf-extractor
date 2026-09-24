"""Exportação dos índices gravados, para conferência ou edição fora da aplicação.

CSV em UTF-8 com datas ISO e decimais com ponto — pensado para curl e scripts.
O XLSX do IGP-DI é o próprio template de importação preenchido, então editar e
reenviar o arquivo exportado é o caminho de edição em massa.
"""

import csv
import io

import openpyxl
from openpyxl.styles import Font

from .importadores import gerar_template_igp_di

COLUNAS_ANP = ("produto", "vigencia_inicio", "vigencia_fim", "regiao", "preco", "origem", "atualizado_em")
COLUNAS_IGP = ("mes", "valor", "origem", "atualizado_em")


def _texto(valor) -> str:
    if valor is None:
        return ""
    if hasattr(valor, "isoformat"):
        return valor.isoformat()
    return str(valor)  # Decimal por str: sem passar por float


def _csv(colunas, linhas) -> bytes:
    saida = io.StringIO()
    escritor = csv.writer(saida, lineterminator="\n")
    escritor.writerow(colunas)
    escritor.writerows([[_texto(v) for v in linha] for linha in linhas])
    return saida.getvalue().encode("utf-8")


def _valores_anp(linha: dict) -> list:
    return [linha[c] for c in COLUNAS_ANP]


def precos_csv(linhas: list[dict]) -> bytes:
    return _csv(COLUNAS_ANP, [_valores_anp(l) for l in linhas])


def precos_xlsx(linhas: list[dict]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Preços ANP"
    ws.append(list(COLUNAS_ANP))
    for celula in ws[1]:
        celula.font = Font(bold=True)
    for linha in linhas:
        # openpyxl não grava datetime com fuso; atualizado_em vai como texto ISO.
        ws.append([
            linha["produto"], linha["vigencia_inicio"], linha["vigencia_fim"], linha["regiao"],
            None if linha["preco"] is None else float(linha["preco"]),
            linha["origem"], _texto(linha["atualizado_em"]),
        ])
        for coluna in (2, 3):
            ws.cell(ws.max_row, coluna).number_format = "dd/mm/yyyy"
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 48
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def indices_csv(linhas: list[dict]) -> bytes:
    return _csv(
        COLUNAS_IGP,
        [[f"{l['mes_ref']:%Y-%m}", l["valor"], l["origem"], l["atualizado_em"]] for l in linhas],
    )


def indices_xlsx(linhas: list[dict]) -> bytes:
    return gerar_template_igp_di(linhas)
